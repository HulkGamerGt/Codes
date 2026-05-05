import openpyxl
from openpyxl.styles import Font, Alignment, numbers
from openpyxl.utils import get_column_letter
import math

wb = openpyxl.Workbook()

# ================================
# Hoja 1: Simulación interactiva
# ================================
ws1 = wb.active
ws1.title = "Simulacion"

# Estilos
titulo_font = Font(bold=True, size=12)
header_font = Font(bold=True)
center = Alignment(horizontal='center')

# Parámetros de entrada
ws1.merge_cells('A1:D1')
ws1['A1'] = "SIMULACIÓN DE LANZAMIENTO DE PROYECTIL"
ws1['A1'].font = titulo_font
ws1['A1'].alignment = center

ws1['A3'] = "Parámetros de entrada"
ws1['A3'].font = Font(bold=True)
ws1['A4'] = "Velocidad inicial v0 (m/s):"
ws1['B4'] = 25.0
ws1['A5'] = "Ángulo θ (°):"
ws1['B5'] = 20.0
ws1['A6'] = "Altura inicial y0 (m):"
ws1['B6'] = 0.0
ws1['A7'] = "Gravedad g (m/s²):"
ws1['B7'] = 9.8

# Cálculos principales (fórmulas)
ws1['A9'] = "Resultados"
ws1['A9'].font = Font(bold=True)
ws1['A10'] = "Tiempo de vuelo tv (s):"
ws1['B10'] = '=2*B4*SIN(RADIANS(B5))/B7'
ws1['A11'] = "Alcance R (m):"
ws1['B11'] = '=B4*COS(RADIANS(B5))*B10'
ws1['A12'] = "Altura máxima ymax (m):"
ws1['B12'] = '=(B4*SIN(RADIANS(B5)))^2/(2*B7)'

# Ajustar ancho de columnas
for col in ['A','B','C','D']:
    ws1.column_dimensions[col].width = 25

# =================================================
# Hoja 2: Cuadro 1 - Resultados para v0=25 m/s
# =================================================
ws2 = wb.create_sheet("Cuadro1_v0_25")
ws2.append(["Ángulo (°)", "Tiempo de vuelo (s)", "Alcance R (m)", "Altura máxima (m)"])
v0 = 25.0
for ang in [20,30,45,60,80]:
    rad = math.radians(ang)
    tv = (2*v0*math.sin(rad))/9.8
    R = v0*math.cos(rad)*tv
    ymax = (v0*math.sin(rad))**2 / (2*9.8)
    ws2.append([ang, round(tv,3), round(R,1), round(ymax,2)])

# Formato
for cell in ws2[1]:
    cell.font = header_font
for col in ws2.columns:
    col_letter = get_column_letter(col[0].column)
    ws2.column_dimensions[col_letter].width = 20

# =================================================
# Hoja 3: Cuadro 2 - Alcance para varias velocidades
# =================================================
ws3 = wb.create_sheet("Cuadro2_Alcances")
ws3.append(["v0 (m/s)"] + [f"{a}°" for a in [20,30,45,60,80]])
for v0 in [25,20,15]:
    fila = [v0]
    for ang in [20,30,45,60,80]:
        rad = math.radians(ang)
        tv = (2*v0*math.sin(rad))/9.8
        R = v0*math.cos(rad)*tv
        fila.append(round(R,1))
    ws3.append(fila)

for cell in ws3[1]:
    cell.font = header_font
for col in ws3.columns:
    col_letter = get_column_letter(col[0].column)
    ws3.column_dimensions[col_letter].width = 12

# =================================================
# Hoja 4: Datos para gráficos (v0=25, θ=20°)
# =================================================
ws4 = wb.create_sheet("Datos_Graficos")
ws4.append(["Tiempo t (s)", "Posición x (m)", "Posición y (m)", "Velocidad vx (m/s)", "Velocidad vy (m/s)"])
v0 = 25.0
theta_deg = 20.0
theta = math.radians(theta_deg)
g = 9.8
tv = 2*v0*math.sin(theta)/g
pts = 6  # 0 a tv, 6 puntos equidistantes
for i in range(pts):
    t = (tv/(pts-1)) * i
    x = v0*math.cos(theta)*t
    y = v0*math.sin(theta)*t - 0.5*g*t**2
    vx = v0*math.cos(theta)
    vy = v0*math.sin(theta) - g*t
    # Redondear a 2 decimales excepto tiempo con 3 decimales
    ws4.append([round(t,3), round(x,2), round(y,2), round(vx,2), round(vy,2)])

for cell in ws4[1]:
    cell.font = header_font
for col in ws4.columns:
    ws4.column_dimensions[get_column_letter(col[0].column)].width = 20

# Guardar archivo
wb.save("Simulacion_Saque_Arquero.xlsx")
print("Archivo 'Simulacion_Saque_Arquero.xlsx' creado exitosamente.")